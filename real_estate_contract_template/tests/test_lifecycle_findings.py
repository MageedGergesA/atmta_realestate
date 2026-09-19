# -*- coding: utf-8 -*-
"""Regressions for the contract-template lifecycle run.

A scripted run generated contracts from a .docx master template and drove the
generated documents through review, approval and signature. Each class below
pins one of the things it found: a lifecycle anybody could drive, a form-version
stamp that rewrote itself, placeholders that failed silently, and an auto-fill
that accepted its own wrong guesses.

The module had no tests at all, so the fixtures live here rather than in a
``common.py``: the estate they need (a project, a unit, a buyer) is already
built by the developer suite's ``DeveloperCommon``.
"""

import base64
import glob
import io
import os
import tempfile

from lxml import etree

from odoo import Command
from odoo.exceptions import AccessError, UserError, ValidationError
from odoo.tests.common import tagged

from odoo.addons.real_estate_developer.tests.common import DeveloperCommon

from docx import Document


# A real 1x1 PNG — res.company.logo is resized on write, so it has to decode.
TINY_PNG = (
    b'iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVR42mP8z8BQDwAEhQGAhKmM'
    b'IQAAAABJRU5ErkJggg=='
)


def make_docx(paragraphs):
    """A .docx whose body is ``paragraphs``, one per line."""
    doc = Document()
    for text in paragraphs:
        doc.add_paragraph(text)
    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()


def docx_text(raw):
    """The visible body text of a rendered .docx."""
    return "\n".join(p.text for p in Document(io.BytesIO(raw)).paragraphs)


def view_arch(env, xmlid):
    return etree.fromstring(env.ref('real_estate_contract_template.' + xmlid).arch)


@tagged('post_install', '-at_install')
class ContractTemplateCase(DeveloperCommon):

    @classmethod
    def setUpClass(cls):
        super().setUpClass()
        cls.Template = cls.env['realestate.contract.template']
        cls.Document = cls.env['realestate.contract.document']
        cls.Wizard = cls.env['realestate.contract.generate.wizard']
        cls.AutoFill = cls.env['realestate.placeholder.autofill.wizard']
        cls.sale_model = cls.env['ir.model']._get('realestate.sale.contract')

        cls.buyer = cls.env['res.partner'].create({
            'name': 'Nadia Farouk', 'zip': '11361', 'city': 'Cairo',
        })
        cls.contract = cls.env['realestate.sale.contract'].create({
            'partner_id': cls.buyer.id,
            'property_id': cls.units[0].id,
            'sale_price': 2500000.0,
        })
        # An employee with no real-estate role at all — the `demo` user of the
        # lifecycle run.
        cls.plain_user = cls.env['res.users'].create({
            'name': 'Plain Employee',
            'login': 'ct_plain_employee',
            'groups_id': [Command.set([cls.env.ref('base.group_user').id])],
        })
        # A leasing agent: the real-estate role every rental role implies, and
        # the tier that prepares documents without being able to approve them.
        cls.agent_user = cls.env['res.users'].create({
            'name': 'Leasing Agent',
            'login': 'ct_leasing_agent',
            'groups_id': [Command.set([
                cls.env.ref('base.group_user').id,
                cls.env.ref('atmta_real_estate.group_realestate_user').id,
            ])],
        })
        cls.template = cls._build_template(cls.env, name='Sale Form EG')

    @classmethod
    def _build_template(cls, env, name='Sale Form', body=None, **vals):
        values = {
            'name': name,
            'kind': 'sale',
            'jurisdiction': 'eg',
            'language': 'bilingual',
            'form_version': 'v1',
            'target_model_id': cls.sale_model.id,
            'docx_template': base64.b64encode(make_docx(body or ['Contract of sale.'])),
            'docx_template_filename': 'form.docx',
        }
        values.update(vals)
        return env['realestate.contract.template'].create(values)

    def _template(self, **kw):
        return self._build_template(self.env, **kw)

    def _wizard(self, template=None, record=None):
        record = record if record is not None else self.contract
        return self.Wizard.create({
            'target_model_name': record._name,
            'record_id': record.id,
            'template_id': (template or self.template).id,
            'output_format': 'docx',
        })

    def _generate(self, **kw):
        wizard = self._wizard(**kw)
        wizard.action_generate()
        return wizard.document_id

    def _make_binding(self, document):
        """Drive ``document`` all the way to Active as a manager."""
        document.action_submit_for_review()
        document.action_approve()
        document.action_send_for_signature()
        document.action_mark_signed()
        document.action_activate()
        return document


class TestDocumentLifecycleSecurity(ContractTemplateCase):
    """B5 — the approval lifecycle was unenforceable."""

    def test_plain_employee_cannot_drive_the_lifecycle(self):
        document = self._generate()
        with self.assertRaises(AccessError):
            document.with_user(self.plain_user).action_submit_for_review()
        self.assertEqual(document.state, 'draft')

    def test_plain_employee_cannot_approve(self):
        document = self._generate()
        document.action_submit_for_review()
        with self.assertRaises(AccessError):
            document.with_user(self.plain_user).action_approve()
        self.assertEqual(document.state, 'under_review')
        self.assertFalse(document.approver_id)

    def test_plain_employee_cannot_activate(self):
        document = self._generate()
        document.action_submit_for_review()
        document.action_approve()
        document.action_mark_signed()
        with self.assertRaises(AccessError):
            document.with_user(self.plain_user).action_activate()
        self.assertEqual(document.state, 'signed')

    def test_signed_document_file_cannot_be_replaced(self):
        document = self._make_binding(self._generate())
        self.assertEqual(document.state, 'active')
        original = document.file_data
        with self.assertRaises(UserError):
            document.write({'file_data': base64.b64encode(b'arbitrary bytes')})
        document.invalidate_recordset()
        self.assertEqual(document.file_data, original)

    def test_signed_document_source_cannot_be_repointed(self):
        document = self._make_binding(self._generate())
        other = self.env['realestate.sale.contract'].create({
            'partner_id': self.buyer.id,
            'property_id': self.units[1].id,
            'sale_price': 100000.0,
        })
        with self.assertRaises(UserError):
            document.write({'target_record_id': other.id})


class TestRealEstateUserTier(ContractTemplateCase):
    """B5, other half — agents prepare, managers approve.

    Closing the hole must not take the ordinary real-estate workflow away with
    it: a Real Estate User drives the preparation steps, and stops at the two
    that make the document binding.
    """

    def test_agent_can_submit_for_review(self):
        document = self._generate()
        document.with_user(self.agent_user).action_submit_for_review()
        self.assertEqual(document.state, 'under_review')

    def test_agent_can_prepare_a_document_for_signature(self):
        document = self._generate()
        document.with_user(self.agent_user).action_submit_for_review()
        document.action_approve()  # the manager's step
        document.with_user(self.agent_user).action_send_for_signature()
        self.assertEqual(document.state, 'sent_for_signature')
        document.with_user(self.agent_user).action_mark_signed()
        self.assertEqual(document.state, 'signed')

    def test_agent_can_cancel_a_draft(self):
        document = self._generate()
        document.with_user(self.agent_user).action_cancel()
        self.assertEqual(document.state, 'cancelled')

    def test_agent_cannot_approve(self):
        document = self._generate()
        document.with_user(self.agent_user).action_submit_for_review()
        with self.assertRaises(AccessError):
            document.with_user(self.agent_user).action_approve()
        self.assertEqual(document.state, 'under_review')
        self.assertFalse(document.approver_id)

    def test_agent_cannot_activate(self):
        document = self._generate()
        document.with_user(self.agent_user).action_submit_for_review()
        document.action_approve()
        document.with_user(self.agent_user).action_mark_signed()
        with self.assertRaises(AccessError):
            document.with_user(self.agent_user).action_activate()
        self.assertEqual(document.state, 'signed')

    def test_agent_cannot_replace_the_file_of_a_signed_document(self):
        document = self._generate()
        document.with_user(self.agent_user).action_submit_for_review()
        document.action_approve()
        document.with_user(self.agent_user).action_mark_signed()
        original = document.file_data
        with self.assertRaises(UserError):
            document.with_user(self.agent_user).write({
                'file_data': base64.b64encode(b'arbitrary bytes'),
            })
        document.invalidate_recordset()
        self.assertEqual(document.file_data, original)

    def test_agent_cannot_replace_the_file_of_an_active_document(self):
        document = self._make_binding(self._generate())
        self.assertEqual(document.state, 'active')
        with self.assertRaises(UserError):
            document.with_user(self.agent_user).write({
                'file_data': base64.b64encode(b'arbitrary bytes'),
            })


class TestFormVersionStamp(ContractTemplateCase):
    """B1 — the form-version stamp followed the template instead of staying put."""

    def test_stamp_survives_a_template_refresh(self):
        document = self._generate()
        self.assertEqual(document.template_version, 'v1')
        self.assertEqual(document.template_kind, 'sale')
        self.assertEqual(document.template_jurisdiction, 'eg')
        self.assertEqual(document.template_language, 'bilingual')

        # The regulator issues a new form: same template record, new file.
        self.template.write({
            'form_version': 'v2',
            'kind': 'rental',
            'jurisdiction': 'sa',
            'language': 'ar',
            'docx_template': base64.b64encode(make_docx(['Revised form.'])),
        })
        document.invalidate_recordset()

        self.assertEqual(document.template_version, 'v1')
        self.assertEqual(document.template_kind, 'sale')
        self.assertEqual(document.template_jurisdiction, 'eg')
        self.assertEqual(document.template_language, 'bilingual')

    def test_next_document_carries_the_new_stamp(self):
        first = self._generate()
        self.template.form_version = 'v2'
        second = self._generate()
        self.assertEqual(first.template_version, 'v1')
        self.assertEqual(second.template_version, 'v2')


class TestPlaceholderRendering(ContractTemplateCase):
    """B4 / B3 — placeholders that fail silently, and "False" in the contract."""

    def test_unknown_placeholder_raises(self):
        template = self._template(
            name='Typo Form',
            body=['Penalty clause: {{ record.no_such_field_at_all }} EGP'],
        )
        with self.assertRaises(UserError):
            template.render(self.contract, output_format='docx')

    def test_known_placeholder_still_renders(self):
        template = self._template(
            name='Good Form', body=['Buyer: {{ record.partner_id.name }}'],
        )
        payload, _filename = template.render(self.contract, output_format='docx')
        self.assertIn('Nadia Farouk', docx_text(payload))

    def test_empty_values_render_empty_not_false(self):
        template = self._template(name='Empty Form', body=[
            'Signing date: {{ record.signing_date }}',
            'Handover date: {{ record.handover_date }}',
            'Cancellation reason: {{ record.cancellation_reason }}',
            'Reservation: {{ record.reservation_id.display_name }}',
            'Reservation record: {{ record.reservation_id }}',
        ])
        text = docx_text(template.render(self.contract, output_format='docx')[0])
        self.assertNotIn('False', text)
        self.assertIn('Signing date:', text)
        self.assertIn('Cancellation reason:', text)

    def test_money_and_dates_can_be_formatted(self):
        template = self._template(name='Money Form', body=[
            'Sale price: {{ record.sale_price|fmt_money }}',
            'Contract date: {{ record.contract_date|fmt_date }}',
            'Signing date: {{ record.signing_date|fmt_date }}',
        ])
        text = docx_text(template.render(self.contract, output_format='docx')[0])
        self.assertIn('2,500,000', text)
        self.assertNotIn('2500000.0', text)
        self.assertNotIn(str(self.contract.contract_date), text)
        self.assertIn('Signing date:', text)


class TestTemplateGuards(ContractTemplateCase):
    """B6 / B7 / B8 / B9 / B11 — the template model's own guards."""

    def test_archived_template_refuses_to_render(self):
        self.template.active = False
        with self.assertRaises(UserError):
            self.template.render(self.contract, output_format='docx')

    def test_archived_template_refuses_through_the_wizard(self):
        wizard = self._wizard()
        self.template.active = False
        with self.assertRaises(UserError):
            wizard.action_generate()

    def test_another_companys_template_is_invisible(self):
        other_company = self.env['res.company'].create({'name': 'Other Co'})
        foreign = self._template(name='Other Co Form', company_id=other_company.id)
        visible = self.Template.with_user(self.plain_user).search([])
        self.assertNotIn(foreign, visible)
        self.assertIn(self.template, visible)

    def test_another_companys_document_is_invisible(self):
        document = self._generate()
        other_company = self.env['res.company'].create({'name': 'Other Co 2'})
        document.sudo().company_id = other_company
        visible = self.Document.with_user(self.plain_user).search([])
        self.assertNotIn(document, visible)

    def test_render_leaves_no_temporary_files(self):
        self.env.company.logo = TINY_PNG
        pattern = os.path.join(tempfile.gettempdir(), 'odoo_company_logo_*')
        before = set(glob.glob(pattern))
        template = self._template(name='Logo Form', body=['Logo: {{ company_logo }}'])
        for _i in range(3):
            template.render(self.contract, output_format='docx')
        self.assertEqual(set(glob.glob(pattern)) - before, set())

    def test_template_file_cannot_be_emptied(self):
        with self.assertRaises(ValidationError):
            self.template.write({'docx_template': False})

    def test_document_file_cannot_be_empty(self):
        with self.assertRaises(ValidationError):
            self.Document.create({
                'name': 'Fileless',
                'template_id': self.template.id,
                'target_model_id': self.sale_model.id,
                'target_record_id': self.contract.id,
                'output_format': 'docx',
                'file_data': False,
                'file_name': 'nothing.docx',
            })

    def test_bad_extension_raises_validation_error(self):
        with self.assertRaises(ValidationError):
            self.template.write({'docx_template_filename': 'form.pdf'})


class TestDeletedSourceRecord(ContractTemplateCase):
    """B10 — the smart button landed on Odoo's "record does not exist" page."""

    def test_open_target_reports_a_deleted_source(self):
        document = self._generate()
        self.contract.unlink()
        with self.assertRaises(UserError):
            document.action_open_target()

    def test_open_target_still_works(self):
        document = self._generate()
        action = document.action_open_target()
        self.assertEqual(action['res_model'], 'realestate.sale.contract')
        self.assertEqual(action['res_id'], self.contract.id)


class TestPlaceholderAutoFill(ContractTemplateCase):
    """B2 — the auto-fill mapped tokens onto the wrong field and pre-accepted them."""

    def _scan(self, tokens):
        template = self._template(
            name='Token Form', body=['%s: [%s]' % (t, t) for t in tokens],
        )
        wizard = self.AutoFill.create({'template_id': template.id})
        wizard.action_scan()
        return wizard

    def test_buyer_token_does_not_suggest_a_postcode(self):
        wizard = self._scan(['Buyer'])
        line = wizard.line_ids.filtered(lambda l: l.token == 'Buyer')
        self.assertTrue(line, "the [Buyer] token was not scanned")
        self.assertNotIn('zip', line.suggested_jinja)

    def test_scan_never_pre_accepts_a_guess(self):
        wizard = self._scan(['Buyer', 'Total Amount', 'Signature'])
        self.assertTrue(wizard.line_ids)
        self.assertFalse(any(wizard.line_ids.mapped('accepted')))
        self.assertFalse(any(wizard.line_ids.mapped('mapped_jinja')))

    def test_nothing_is_applied_without_an_explicit_accept(self):
        wizard = self._scan(['Buyer'])
        with self.assertRaises(UserError):
            wizard.action_apply()


class TestGenerateWizardUx(ContractTemplateCase):
    """The wizard's own success screen, and the document views."""

    def test_generate_stays_in_the_dialog(self):
        wizard = self._wizard()
        action = wizard.action_generate()
        self.assertEqual(action['type'], 'ir.actions.act_window')
        self.assertEqual(action['res_model'], wizard._name)
        self.assertEqual(action['res_id'], wizard.id)
        self.assertTrue(wizard.generated)
        self.assertTrue(wizard.document_id)
        self.assertTrue(wizard.output_file)

    def test_document_views_forbid_manual_creation(self):
        for xmlid in ('view_contract_document_list', 'view_contract_document_form'):
            self.assertEqual(view_arch(self.env, xmlid).get('create'), '0', xmlid)

    def test_download_button_is_hidden_without_a_file(self):
        arch = view_arch(self.env, 'view_contract_document_form')
        button = arch.xpath("//button[@name='action_download']")
        self.assertTrue(button)
        self.assertTrue(button[0].get('invisible'),
                        "Download is offered on a document with no file")

    def test_file_size_is_not_shown_as_a_float(self):
        arch = view_arch(self.env, 'view_contract_document_form')
        node = arch.xpath("//field[@name='file_size']")[0]
        self.assertNotEqual(node.get('widget'), 'float')
